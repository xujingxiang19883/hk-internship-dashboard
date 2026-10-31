# -*- coding: utf-8 -*-
"""Round-2 new US + HK QR/QT targets, cross-checked against the 94-entry applied list
from the 2026 dashboard snapshot (2026-10-31)."""
import sys, os
XLSX_SKILL_DIR = r"C:\Users\HUAWEI\.zcode\cli\plugins\cache\zcode-plugins-official\spreadsheets\0.1.7\skills\xlsx"
for sub in [XLSX_SKILL_DIR, os.path.join(XLSX_SKILL_DIR, "templates")]:
    if sub not in sys.path:
        sys.path.insert(0, sub)

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

US = [
 ("Bridgewater Associates", "2027 Investment Associate Intern", "New York",
  "Open (4周前发布)", "https://job-boards.greenhouse.io/bridgewater89/jobs/8395041002",
  "研究型投资助理实习，经济/数据驱动研究，数学要求高，未投过"),
 ("Citi", "Markets – Quantitative Analysis, Summer Analyst (NYC, 2027)", "New York",
  "Open", "https://citi.wd5.myworkdayjobs.com/2/job/New-York-New-York-United-States/Markets---Quantitative-Analysis--Summer-Analyst---New-York-City---US--2027_25928747",
  "Citi 美国站首次申请（此前只讨论过 HK 岗）"),
 ("Graham Capital Management", "2027 Graham Capital Summer Internship – CT / NYC", "New York / Norwalk, CT",
  "Open", "https://boards.greenhouse.io/grahamcapitalmanagement/jobs/4721074005?gh_jid=4721074005",
  "系统化宏观+自主策略基金，QR/QT 混合培养"),
 ("Graviton Research Capital", "Quantitative Researcher – 2027 Summer Internship", "Amsterdam / London / New York / Singapore",
  "Open（1周前发布）", "https://job-boards.greenhouse.io/gravitonresearchcapital/jobs/8899673002",
  "高频交易公司 QR 实习"),
 ("Man Group (AHL)", "Systematic Fast Trading Intern", "New York",
  "Open（4天前发布）", "https://job-boards.eu.greenhouse.io/mangroup/jobs/5012801101",
  "AHL 系统化快交易部门，QT 方向"),
 ("Quadeye", "Intern: Quant Researcher", "New York",
  "Open", "https://www.quadeye.com/jobs/199893000001828092/intern-quant-researcher?source=CareerSite",
  "高频自营，QR 实习"),
 ("Quadrature Capital", "Quadrature Capital Internships", "London / New York",
  "Open", "https://job-boards.greenhouse.io/quadraturecapital/jobs/4255974",
  "系统化交易研究实习"),
 ("Tudor Investment Corporation", "Tudor Internships 2027", "NY / Stamford / London / Singapore",
  "Open（2天前发布）", "https://job-boards.greenhouse.io/tudorgroup/jobs/8103095",
  "宏观交易，全球轮岗实习"),
 ("Capstone Investment Advisors", "Summer 2027 – Investment Internship NYC", "New York",
  "Open", "https://job-boards.greenhouse.io/capstoneinvestmentadvisors/jobs/8827973002",
  "波动率/尾部风险策略，QR 方向"),
 ("American Century Investments", "Quantitative Research Intern + Systematic Credit Intern", "Kansas City / New York",
  "Open", "https://americancentury.wd5.myworkdayjobs.com/AmericanCenturyInvestments/job/Kansas-City-Missouri/Quantitative-Research-Intern_R0005728",
  "量化资管；QR 与系统化信用两个口（Systematic Credit: R0005759，在 NY）"),
 ("GSA Capital", "Quantitative Researcher – Intern", "London / New York",
  "Open", "https://job-boards.greenhouse.io/embed/job_app?for=gsacapital&token=8570661002",
  "伦敦起家的量化自营，QR 实习"),
 ("Balyasny Asset Management", "Quantitative Analyst – Macro Investment Team (Summer Internship)", "New York",
  "Open", "https://bambusdev.my.site.com/s/details?jobReq=Quantitative-Analyst---Macro-Investment-Team--Summer-Internship-_REQ8453",
  "10/29 已投 BAAM 两个 QR 岗（Systematic Multi-Asset、Risk）；此为宏观 QR，可加投"),
 ("Brevan Howard", "2027 Summer Internship – Systematic Trading / Systematic Data Strategy", "New York",
  "Open", "https://brevanhoward.wd3.myworkdayjobs.com/BH_ExternalCareers/job/New-York/XMLNAME-2027-Summer-Internship-Program---Systematic-Trading--New-York_JR101591",
  "已投 AI & Quantitative Analyst track；系统化交易(JR101591)/数据策略(JR101596)为不同 track，可加投"),
 ("DV Trading", "Quantitative Trading Intern – Summer 2027 (DV Equities)", "New York",
  "Open", "https://job-boards.greenhouse.io/dvtrading/jobs/4702083005",
  "已投 DV QR Equities；此为 QT 方向，可加投"),
]

HK = [
 ("Tower Research Capital", "Quantitative Researcher Intern, Bachelor's or Master's", "Singapore / Hong Kong / Shanghai / Sydney",
  "Open", "https://www.tower-research.com/open-positions/?gh_jid=8168750",
  "按此前指示：美国已投（QT/R），香港 QR 不冲突可加投"),
 ("Jump Trading", "Campus Quantitative Trader (Intern)", "Hong Kong / Shanghai",
  "Open", "https://www.jumptrading.com/hr/job?gh_jid=8027922",
  "已投 HK QR Intern + US Campus QR；此为 QT 岗位，可加投"),
 ("Century Frontier (世纪前沿)", "量化研究员（应届/实习）Quant Researcher", "上海 / 深圳 / 北京 / 香港",
  "Open", "https://app.mokahr.com/m/campus_apply/centuryfrontier/24842?hash=%23%2Fjob%2F7fdca733-f830-4f90-b37b-86b6aa163857",
  "中资量化，未投过（已投的五家为明汯/幻方/九坤/灵均/宽德）"),
 ("Bastion Trading", "Internship for Trading & Quant (Summer 2027)", "Hong Kong",
  "Open", "https://bastiontrading.com/careers/apply",
  "HK 加密做市商，QR/QT 实习，与 Polymarket 经历契合"),
 ("Dymon Asia Capital", "2027 AI Equity Long/Short (ELS) Fellowship", "Hong Kong",
  "⚠️ 表单要求 2027年7月起 6 个月全职实习（与 MFE 期末学期冲突，同 UBS）— 确认能承诺再投", "https://job-boards.anz.greenhouse.io/dymonasiacapitalsingaporepteltd/jobs/4006743201",
  "徐曾在 Dymon Asia 实习（老东家）；但 6 个月全职要求未解决前不建议投"),
 ("TSY Capital", "Quantitative Researcher Intern (2026/2027)", "Hong Kong",
  "Open", "https://tsycapital.com/career.html",
  "HK 量化基金 QR 实习"),
 ("XY Capital", "2027 Summer Internship – Quant Research and Trading", "Hong Kong",
  "⚠️ JD 页标注截止 2026-10-16 已过 — 门户或仍接受，可一试", "https://xycapital.careers.hibob.com/jobs/ca79ec51-06d9-4935-988e-039adf65db51",
  "HK 量化基金 QR/QT 夏季实习；CV 已备好"),
]

NAVY = "1B2A4A"; N100 = "F7F7F5"; N900 = "37352F"; WHITE = "FFFFFF"

def make_sheet(ws, title, rows):
    ws.merge_cells("A1:G1")
    c = ws["A1"]; c.value = title
    c.font = Font(bold=True, size=13, color=NAVY)
    ws.row_dimensions[1].height = 24
    headers = ["#", "公司", "职位", "地点", "截止 / 状态", "申请入口", "备注"]
    HR = 3
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=HR, column=col, value=h)
        cell.font = Font(bold=True, size=10, color=WHITE)
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.alignment = Alignment(horizontal="left", vertical="center")
    for idx, (comp, role, loc, dl, url, note) in enumerate(rows, 1):
        r = HR + idx
        fill = PatternFill("solid", fgColor=N100) if idx % 2 == 1 else None
        for col, v in enumerate([idx, comp, role, loc, dl, url, note], 1):
            cell = ws.cell(row=r, column=col, value=v)
            cell.font = Font(size=9.5, color=N900)
            cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=(col in (2, 3, 7)))
            if fill: cell.fill = fill
        lc = ws.cell(row=r, column=6)
        lc.hyperlink = url
        lc.value = "点击申请"
        lc.font = Font(size=9, color="0B6D8A", underline="single")
        lc.alignment = Alignment(horizontal="left", vertical="top")
        ws.row_dimensions[r].height = 34
    widths = [4, 26, 46, 26, 16, 12, 52]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A4"

wb = Workbook()
ws1 = wb.active; ws1.title = "美国新目标"
make_sheet(ws1, "美国 QR/QT 实习新目标 · 2026-10-31 · 已对照 dashboard 全部 94 条申请记录去重", US)
ws2 = wb.create_sheet("香港新目标")
make_sheet(ws2, "香港 QR/QT 实习新目标 · 2026-10-31 · 已对照 dashboard 去重（全职graduate岗与受限公司已剔除）", HK)
wb.properties.creator = "Z.ai"
out = r"D:\zcode\hk-dashboard\New_Targets_Round2_US_HK_v2.xlsx"
wb.save(out)
print("saved:", out, "| US rows:", len(US), "| HK rows:", len(HK))
